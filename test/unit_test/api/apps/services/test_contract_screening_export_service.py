#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
import openpyxl
import pytest
from docx import Document

from api.apps.services.contract_screening_service import ContractScreeningError
from api.apps.services.contract_screening_export_service import create_screening_export


class FakeRepository:
    def __init__(self, payload):
        self.payload = payload
        self.exports = []
        self.updated = []
        self.build_calls = []

    def build_results_payload(self, tenant_id, task_id, user_id=None):
        assert tenant_id == "tenant-1"
        assert task_id == "task-1"
        self.build_calls.append((tenant_id, task_id, user_id))
        return self.payload

    class ContractScreeningExportService:
        pass

    def create_export_record(self, record):
        self.exports.append(record)
        return record

    def mark_export_done(self, export_id, file_name, file_key):
        self.updated.append((export_id, file_name, file_key))


def _payload(status="done"):
    return {
        "task_id": "task-1",
        "prompt": "筛选付款周期超过60天的合同",
        "strategy": {
            "conditions": [{"id": "payment_terms", "label": "付款周期", "keywords": ["付款"], "enabled": True}],
        },
        "status": status,
        "skipped": {"unparsed": 1},
        "items": [
            {
                "id": "result-1",
                "contract_id": "doc-1",
                "name": "采购合同.pdf",
                "status": "matched",
                "risk": "高",
                "score": 91,
                "reason": "付款周期90天。",
                "actions": ["请求法务复核"],
                "evidence": [
                    {
                        "source": "合同正文",
                        "ref": "第3页 / chunk-1",
                        "page": 3,
                        "chunk_id": "chunk-1",
                        "text": "付款期限为验收合格后90日内。",
                        "condition_id": "payment_terms",
                    }
                ],
            }
        ],
    }


def test_create_screening_export_rejects_unfinished_tasks(tmp_path):
    with pytest.raises(ContractScreeningError) as exc:
        create_screening_export(
            tenant_id="tenant-1",
            user_id="user-1",
            task_id="task-1",
            export_format="excel",
            output_dir=tmp_path,
            repository=FakeRepository(_payload(status="running")),
        )

    assert "Only completed screening tasks can be exported" in exc.value.message


def test_create_screening_export_writes_excel_sheets(tmp_path):
    repository = FakeRepository(_payload())

    result = create_screening_export(
        tenant_id="tenant-1",
        user_id="user-1",
        task_id="task-1",
        export_format="excel",
        output_dir=tmp_path,
        repository=repository,
    )

    workbook = openpyxl.load_workbook(result["file_key"])
    assert workbook.sheetnames == ["筛选摘要", "合同结果", "证据明细"]
    assert workbook["筛选摘要"]["B1"].value == "筛选付款周期超过60天的合同"
    assert workbook["合同结果"]["A2"].value == "采购合同.pdf"
    assert workbook["证据明细"]["F2"].value == "付款期限为验收合格后90日内。"
    assert repository.updated[0][0] == result["export_id"]
    assert repository.build_calls == [("tenant-1", "task-1", "user-1")]


def test_create_screening_export_writes_word_document(tmp_path):
    result = create_screening_export(
        tenant_id="tenant-1",
        user_id="user-1",
        task_id="task-1",
        export_format="word",
        output_dir=tmp_path,
        repository=FakeRepository(_payload()),
    )

    document = Document(result["file_key"])
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "筛选付款周期超过60天的合同" in text
    assert "采购合同.pdf" in text
    assert "付款期限为验收合格后90日内。" in text
